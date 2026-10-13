"""네이버 블로그 실제 화면 수집기, 클라우드 웹버전.

단일 Railway 서비스/단일 worker 기준. 실제 수집은 Playwright Chromium에서
화면 DOM을 읽으며 네이버 검색 API를 사용하지 않습니다.
"""
from __future__ import annotations

import os
import secrets
import shutil
import threading
import time
import uuid
from collections import deque
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, HTTPException, Header, Request
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, field_validator

from screen_crawler import run_batch, split_keywords

ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("CRAWLER_DATA_DIR", "/tmp/naver-crawler-jobs"))
MAX_JOBS = 15
JOB_TTL_S = 2 * 60 * 60
ACCESS_KEY = os.environ.get("APP_ACCESS_KEY", "")
REQUIRE_KEY = bool(os.environ.get("RAILWAY_ENVIRONMENT")) or bool(ACCESS_KEY)
SINGLE_TASK = ThreadPoolExecutor(max_workers=1, thread_name_prefix="naver-playwright")
LOCK = threading.RLock()
JOBS: dict[str, dict] = {}
ACTIVE: set[str] = set()

app = FastAPI(title="네이버 블로그 실제 화면 수집기", docs_url=None, redoc_url=None)
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")


class JobRequest(BaseModel):
    keywords: str = Field(min_length=1, max_length=650)
    device: str = "mobile"
    limit: int = Field(default=20, ge=1, le=100)
    scrolls: int = Field(default=15, ge=0, le=35)
    pause: float = Field(default=1.5, ge=1.0, le=15.0)
    lookup_missing_names: bool = True
    screenshots: bool = False

    @field_validator("device")
    @classmethod
    def check_device(cls, v):
        if v not in {"mobile", "desktop"}:
            raise ValueError("모바일 또는 PC만 선택할 수 있습니다.")
        return v

    @field_validator("keywords")
    @classmethod
    def check_keywords(cls, v):
        words = split_keywords(v)
        if not words or len(words) > 8 or any(len(w) > 65 for w in words):
            raise ValueError("키워드는 최대 8개, 키워드당 65자 이하로 입력해주세요.")
        return v


def authorize(key: str | None):
    if REQUIRE_KEY and not ACCESS_KEY:
        raise HTTPException(503, "관리자가 서버 접근코드를 설정해야 합니다.")
    if ACCESS_KEY and not secrets.compare_digest(key or "", ACCESS_KEY):
        raise HTTPException(401, "접근코드를 확인해주세요.")


def clean_old_jobs():
    now = time.time()
    with LOCK:
        for jid, job in list(JOBS.items()):
            if job["status"] in ("완료", "오류", "중지", "부분 실패") and now - job["created_ts"] > JOB_TTL_S:
                JOBS.pop(jid)
                shutil.rmtree(DATA_DIR / jid, ignore_errors=True)
        while len(JOBS) >= MAX_JOBS:
            ended = sorted((j for j in JOBS.values() if j["status"] in ("완료", "오류", "중지", "부분 실패")), key=lambda j: j["created_ts"])
            if not ended:
                break
            victim = ended[0]["id"]
            del JOBS[victim]
            shutil.rmtree(DATA_DIR / victim, ignore_errors=True)


def job_public(job: dict):
    return {"id": job["id"], "status": job["status"], "created": job["created"],
            "finished": job.get("finished"), "count": job["count"], "errors": job["errors"],
            "logs": list(job["logs"]), "downloadReady": bool(job["output"].exists()),
            "error": job.get("error", ""), "screenshots": job.get("screenshots", []), "rows": job.get("rows", [])}


def run_job(jobid: str, cfg: dict):
    with LOCK:
        job = JOBS[jobid]
        job["status"] = "수집 중"

    def log(message):
        with LOCK:
            if jobid in JOBS:
                JOBS[jobid]["logs"].append(f"[{datetime.now():%H:%M:%S}] {message}")

    try:
        count, errors, debug_dir = run_batch(cfg, job["stop"], log)
        # 웹 표에는 실제 생성된 엑셀 데이터만 표시합니다.
        from openpyxl import load_workbook
        items = []
        if job["output"].is_file():
            wb = load_workbook(job["output"], read_only=True, data_only=True)
            try:
                ws = wb["화면 수집 결과"]
                rows_iter = ws.iter_rows(values_only=True)
                keys = next(rows_iter)
                for vals in rows_iter:
                    items.append(dict(zip(keys, vals)))
            finally:
                wb.close()
        with LOCK:
            job["count"] = count
            job["rows"] = items[:800]
            job["errors"] = errors
            job["screenshots"] = [p.name for p in debug_dir.glob("*.png")][:15] if debug_dir and debug_dir.exists() else []
            # 키워드 일부에서 오류가 나면 완료로 위장하지 않고 일부 실패를 명시
            if job["stop"].is_set():
                job["status"] = "중지"
            elif errors:
                job["status"] = "부분 실패" if count else "오류"
            else:
                job["status"] = "완료"
            if job["status"] == "부분 실패":
                job["error"] = f"일부 키워드({errors}개) 수집 실패. 엑셀의 키워드별 현황을 확인하세요."
            elif job["status"] == "오류":
                job["error"] = "모든 키워드 수집에 실패했습니다. 진행 기록을 확인해주세요."
        log(f"작업 종료. 수집 {count}건 / 실패 키워드 {errors}개")
    except Exception as ex:
        log(f"서버 실행 오류: {type(ex).__name__}: {str(ex)[:250]}")
        with LOCK:
            job["status"] = "오류"
            job["error"] = f"{type(ex).__name__}: {str(ex)[:200]}"
    finally:
        with LOCK:
            job["finished"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ACTIVE.discard(jobid)


@app.get("/health")
def health():
    return {"ok": True, "app": "naver-blog-screen-crawler", "authConfigured": bool(ACCESS_KEY) or not REQUIRE_KEY}


@app.get("/")
def home():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/api/browser-check")
def browser_check(x_access_key: str | None = Header(default=None)):
    """서버의 Chromium 실행과 모의 검색화면 DOM 추출을 확인합니다."""
    authorize(x_access_key)
    with LOCK:
        if ACTIVE:
            raise HTTPException(409, "다른 수집이 진행 중이라 브라우저 진단을 시작할 수 없습니다.")
    from playwright.sync_api import sync_playwright
    from screen_crawler import JS_FILE
    html = '''<main id="main_pack"><article class="result_item"><a class="title_link" href="https://m.blog.naver.com/checkuser/1234567890">브라우저 검사 게시글</a><a href="https://m.blog.naver.com/checkuser">테스트 블로그</a></article></main>'''
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True, executable_path=os.environ.get("BROWSER_EXECUTABLE_PATH") or None, args=["--no-sandbox"])
            try:
                page = browser.new_page()
                page.set_content(html)
                rows = page.evaluate(JS_FILE.read_text(encoding="utf-8"))
            finally:
                browser.close()
        if not (len(rows) == 1 and rows[0]["blogName"] == "테스트 블로그"):
            raise RuntimeError("검색화면 추출 결과를 확인할 수 없습니다.")
        return {"browserReady": True, "domExtractionReady": True, "liveNaverVerified": False}
    except Exception as ex:
        raise HTTPException(503, f"Chromium 진단 실패: {type(ex).__name__}: {str(ex)[:160]}")


@app.post("/api/jobs", status_code=202)
def create_job(payload: JobRequest, x_access_key: str | None = Header(default=None)):
    authorize(x_access_key)
    clean_old_jobs()
    with LOCK:
        if ACTIVE:
            raise HTTPException(409, "다른 수집이 진행 중입니다. 완료 후 다시 시도해주세요.")
        if len(JOBS) >= MAX_JOBS:
            raise HTTPException(429, "결과 저장 한도에 도달했습니다. 잠시 후 다시 시도해주세요.")
        jobid = uuid.uuid4().hex
        outdir = DATA_DIR / jobid
        outdir.mkdir(parents=True, exist_ok=True)
        output = outdir / f"네이버_블로그_수집_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        cfg = {"keywords": split_keywords(payload.keywords), "device": payload.device,
               "limit": payload.limit, "scrolls": payload.scrolls, "pause": payload.pause,
               "lookup_missing_names": payload.lookup_missing_names,
               "screenshots": payload.screenshots,
               "headless": True, "dest": str(output)}
        JOBS[jobid] = {"id": jobid, "status": "대기", "created": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                       "created_ts": time.time(), "finished": None, "count": 0, "errors": 0,
                       "logs": deque(["수집 작업이 생성되었습니다."], maxlen=220), "stop": threading.Event(),
                       "output": output, "screenshots": [], "rows": [], "error": ""}
        ACTIVE.add(jobid)
    SINGLE_TASK.submit(run_job, jobid, cfg)
    return {"id": jobid, "status": "대기"}


@app.get("/api/jobs/{jobid}")
def get_job(jobid: str, x_access_key: str | None = Header(default=None)):
    authorize(x_access_key)
    with LOCK:
        if jobid not in JOBS:
            raise HTTPException(404, "작업이 없거나 보관 시간이 만료되었습니다.")
        return job_public(JOBS[jobid])


@app.post("/api/jobs/{jobid}/stop")
def stop_job(jobid: str, x_access_key: str | None = Header(default=None)):
    authorize(x_access_key)
    with LOCK:
        if jobid not in JOBS:
            raise HTTPException(404, "작업을 찾을 수 없습니다.")
        JOBS[jobid]["stop"].set()
        JOBS[jobid]["logs"].append("중지 요청을 받았습니다. 현재 단계 종료 후 저장합니다.")
        return {"ok": True}


@app.get("/api/jobs/{jobid}/download")
def download_job(jobid: str, x_access_key: str | None = Header(default=None)):
    authorize(x_access_key)
    with LOCK:
        job = JOBS.get(jobid)
        if not job or not job["output"].is_file():
            raise HTTPException(404, "다운로드할 엑셀이 없습니다.")
        if job["status"] not in ("완료", "부분 실패", "오류", "중지"):
            raise HTTPException(409, "수집이 진행 중입니다.")
        path = job["output"]
    return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=path.name)


@app.get("/api/jobs/{jobid}/screenshots/{name}")
def download_screenshot(jobid: str, name: str, x_access_key: str | None = Header(default=None)):
    authorize(x_access_key)
    with LOCK:
        job = JOBS.get(jobid)
        if not job or name not in job["screenshots"]:
            raise HTTPException(404, "화면 파일이 없습니다.")
        candidates = list((DATA_DIR / jobid).glob(f"네이버_수집_진단화면_*/*"))
        found = next((p for p in candidates if p.name == name and p.suffix.lower() == ".png"), None)
    if not found:
        raise HTTPException(404, "화면 파일을 찾을 수 없습니다.")
    return FileResponse(found, media_type="image/png", filename=name)
